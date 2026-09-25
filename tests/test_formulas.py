import hashlib
import json
import re
import time
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from xlsprint.formulas import (
    BUILTIN_FUNCTIONS,
    SINGLE_THREADED_FUNCTIONS,
    VOLATILE_FUNCTIONS,
    _classify,
    compact_areas,
    fingerprint,
    inspect_workbook,
    inspect_workbook_and_plan,
    inspect_workbook_for_plan,
    normalize_formula,
    rect_a1,
    redact_name,
    tokenize,
    write_formulas_json,
)
from xlsprint.synthbook import make_synthetic_workbook

SALT = "test-salt"


def classify(formula, sheet="Host", names=frozenset()):
    return _classify(tokenize(formula), sheet, names)


def fp(formula, row, col):
    return fingerprint(normalize_formula(formula, row, col))


@pytest.fixture
def synth(tmp_path):
    return make_synthetic_workbook(tmp_path / "synth.xlsx", seed=3)


@pytest.fixture
def synth_out(synth):
    return inspect_workbook(synth, names="clear")


def group_at(out, sheet, area):
    matches = [g for g in out["groups"] if g["sheet"] == sheet and area in g["areas"]]
    assert len(matches) == 1, (sheet, area, [g["areas"] for g in out["groups"]])
    return matches[0]


# --- normalization ---------------------------------------------------------

def test_relative_copies_share_fingerprint():
    assert fp("=A1+1", 1, 2) == fp("=A2+1", 2, 2)
    assert normalize_formula("=A1+1", 1, 2) == "RC[-1]+<N>"


def test_absolute_and_relative_differ():
    assert fp("=$A$1+1", 1, 2) != fp("=A1+1", 1, 2)
    assert fp("=$A$1+1", 1, 2) == fp("=$A$1+1", 7, 2)
    # mixed anchors
    assert normalize_formula("=$A1+A$1", 3, 3) == "R[-2]C1+R1C[-2]"


def test_period_to_date_pattern_groups():
    assert fp("=SUM($A$1:$A1)", 1, 2) == fp("=SUM($A$1:$A50)", 50, 2)


def test_whole_column_row_and_sheet_refs():
    assert normalize_formula("=SUM(A:A)", 5, 3) == "SUM(C[-2]:C[-2])"
    assert normalize_formula("=SUM($1:$2)", 5, 3) == "SUM(R1:R2)"
    n = normalize_formula("='My Sheet'!B2+[1]Other!$C$3", 2, 2)
    assert n == "S:MY SHEET!RC+S:[1]OTHER!R3C3"
    assert fp("='My Sheet'!B2", 2, 2) != fp("='Their Sheet'!B2", 2, 2)


def test_literal_stripping_whitespace_and_case():
    assert fp("=A1*5", 1, 2) == fp("=A1*7.25E+3", 1, 2)
    assert fp('=A1&"alpha"', 1, 2) == fp('=A1&"be""ta"', 1, 2)
    assert fp("=sum( A1 , 2 )", 1, 2) == fp("=SUM(A1,3)", 1, 2)
    n = normalize_formula('=IF(A1>12345,"secret",{1,2;3,4})', 1, 2)
    assert "12345" not in n and "secret" not in n and "<N>" in n and "<S>" in n


def test_xlfn_prefix_stripped():
    assert fp("=_xlfn.STDEV.S(A1:A3)", 1, 2) == fp("=STDEV.S(A1:A3)", 1, 2)
    f = classify("=_xlfn._xlws.SORT(A1:A3)+_xlfn.STDEV.S(A1:A3)")
    assert f.functions == {"SORT", "STDEV.S"}
    assert not f.udfs
    assert f.dynamic_array


def test_tokenizer_is_total():
    for junk in ["=", "=(((", "=\"unterminated", "=A1##", "={1,2", "=SUM(A1", "=)", "=Sheet1!#REF!+1",
                 "=A1:INDEX(B:B,3)", "=A1:A5 B1:B5", "=@A1:A3", "=A1#", "=_xlpm.x", "=été+1"]:
        tokenize(junk)
        normalize_formula(junk, 1, 1)
        classify(junk)


def test_intersection_space_matters():
    assert fp("=A1:A5 B1:B5", 1, 3) != fp("=A1:A5B1:B5", 1, 3)
    assert fp("=A1 + B1", 1, 3) == fp("=A1+B1", 1, 3)


# --- detectors ---------------------------------------------------------------

def test_function_sets():
    assert VOLATILE_FUNCTIONS == {"RAND", "RANDBETWEEN", "RANDARRAY", "NOW", "TODAY",
                                  "OFFSET", "INDIRECT", "CELL", "INFO"}
    assert VOLATILE_FUNCTIONS <= BUILTIN_FUNCTIONS
    assert SINGLE_THREADED_FUNCTIONS <= BUILTIN_FUNCTIONS
    for f in ["XLOOKUP", "LET", "LAMBDA", "TEXTSPLIT", "VSTACK", "SUMIFS", "NETWORKDAYS.INTL"]:
        assert f in BUILTIN_FUNCTIONS


@pytest.mark.parametrize("formula,expected", [
    ("=RAND()", {"RAND"}), ("=NOW()-TODAY()", {"NOW", "TODAY"}),
    ("=OFFSET(A1,1,1)", {"OFFSET"}), ('=INDIRECT("A1")', {"INDIRECT"}),
    ('=INFO("osversion")', {"INFO"}), ("=_xlfn.RANDARRAY(3)", {"RANDARRAY"}),
    ("=RANDBETWEEN(1,6)", {"RANDBETWEEN"}), ("=SUM(A1:A3)", set()),
])
def test_volatile(formula, expected):
    assert classify(formula).volatile == expected


@pytest.mark.parametrize("formula,single", [
    ('=CELL("address",A1)', True),
    ('=CELL("Format",A1)', True),
    ('=CELL( "format" ,A1)', True),
    ('=CELL("width",A1)', False),
    ("=CELL(B1,A1)", False),          # info_type not a literal: not decidable
    ('=CELL("row")', False),
    ('=ADDRESS(1,1,1,TRUE,"Sheet")', True),
    ("=ADDRESS(1,1,1,TRUE,B2)", True),
    ("=ADDRESS(1,1,1,TRUE,)", False),  # empty sheet_text
    ("=ADDRESS(1,1)", False),
    ("=ADDRESS(1,1,1,TRUE)", False),
    ('=ADDRESS(ROW(A1),MAX(1,2),1,TRUE,"S")', True),
    ('=ADDRESS(1,CELL("width",A1))', False),
])
def test_cell_address_rules(formula, single):
    f = classify(formula)
    flagged = {"CELL", "ADDRESS"} & f.single_threaded
    assert bool(flagged) == single, (formula, f.single_threaded)


@pytest.mark.parametrize("formula,name", [
    ("=PHONETIC(A1)", "PHONETIC"), ('=INDIRECT("A1")', "INDIRECT"),
    ('=GETPIVOTDATA("x",A1)', "GETPIVOTDATA"), ('=CUBEVALUE("c")', "CUBEVALUE"),
    ("=ERROR.TYPE(A1)", "ERROR.TYPE"), ('=HYPERLINK("x")', "HYPERLINK"),
])
def test_single_threaded(formula, name):
    assert name in classify(formula).single_threaded


def test_udfs():
    f = classify("=MYUDF(A1)+_xludf.OTHER(1)+SUM(A1)+_xll.ADDIN(2)")
    assert f.udfs == {"MYUDF", "OTHER", "ADDIN"}
    assert f.functions == {"SUM"}
    assert "UDF" in f.single_threaded
    # a call to a defined name (LAMBDA) is not a UDF
    g = classify("=MyLambda(A1)", names=frozenset({"MYLAMBDA"}))
    assert not g.udfs and g.lambda_calls == {"MYLAMBDA"}


def test_ref_flags():
    f = classify("=Other!A1+'My Sheet'!B2+A1", sheet="Host")
    assert f.cross_sheet and f.references == 3 and not f.external_ref
    assert not classify("=Host!A1+A2", sheet="Host").cross_sheet
    assert classify("=Sheet1:Sheet3!A1").cross_sheet
    assert classify("=[1]Sheet1!A1").external_ref
    assert classify("=[2]!RemoteName").external_ref
    assert classify("=SUM(A:A)").whole_column_ref
    assert classify("=SUM($3:$4)").whole_row_ref
    assert not classify("=SUM(A1:A9)").whole_column_ref
    assert classify("=A1#").dynamic_array
    assert classify("=_xlfn.ANCHORARRAY(A1)").dynamic_array


# --- areas -------------------------------------------------------------------

def test_compact_areas():
    col = [(r, 2) for r in range(2, 102)]
    assert [rect_a1(r) for r in compact_areas(col)] == ["B2:B101"]
    block = [(r, c) for r in range(1, 11) for c in range(3, 7)]
    assert [rect_a1(r) for r in compact_areas(block)] == ["C1:F10"]
    row = [(4, c) for c in range(1, 30)]
    assert [rect_a1(r) for r in compact_areas(row)] == ["A4:AC4"]
    ell = [(r, 1) for r in range(1, 6)] + [(5, c) for c in range(2, 6)]
    assert len(compact_areas(ell)) == 2
    disjoint = [(1, 1), (3, 1), (1, 3)]
    assert sorted(rect_a1(r) for r in compact_areas(disjoint)) == ["A1", "A3", "C1"]
    assert compact_areas([(1, 1), (1, 1)]) == [(1, 1, 1, 1)]


def test_area_cap(tmp_path):
    wb = Workbook()
    ws = wb.active
    for i in range(40):
        r = 1 + 2 * i
        ws.cell(row=r, column=2, value="=A%d*2" % r)
    p = tmp_path / "cap.xlsx"
    wb.save(p)
    out = inspect_workbook(p, names="clear")
    (g,) = out["groups"]
    assert g["cells"] == 40 and g["areas_total"] == 40 and len(g["areas"]) == 32
    plan = inspect_workbook_for_plan(p)
    assert len(plan["groups"][0]["areas"]) == 40


# --- workbook-level ----------------------------------------------------------

def test_synth_grouping(synth_out):
    out = synth_out
    assert out["schema"] == "xlsprint.formulas/1"
    assert len(out["workbook_sha256_prefix"]) == 10
    sheets = {s["sheet"]: s for s in out["sheets"]}
    assert sheets["Inputs"]["formula_cells"] == 0
    assert sheets["Report"]["formula_cells"] == 3
    assert group_at(out, "Model", "A2:A201")["cells"] == 200
    assert group_at(out, "Model", "B2:B201")["functions"] == ["SUM"]
    assert group_at(out, "Model", "C3:C201")["cells"] == 199
    assert group_at(out, "Model", "C2")["cells"] == 1
    assert out["totals"]["formula_cells"] == sum(g["cells"] for g in out["groups"])
    assert out["totals"]["formula_cells"] == sum(s["formula_cells"] for s in out["sheets"])
    assert out["totals"]["groups"] == len(out["groups"])
    ids = [g["group"] for g in out["groups"]]
    assert ids == ["G%04d" % (i + 1) for i in range(len(ids))]
    assert out["limitations"]


def test_synth_detectors(synth_out):
    out = synth_out
    a = group_at(out, "Model", "A2:A201")
    assert a["cross_sheet"] and not a["volatile"] and a["references"] == 2
    d = group_at(out, "Model", "D2:D201")
    assert d["volatile"] and d["volatile_functions"] == ["OFFSET"] and not d["single_threaded"]
    e = group_at(out, "Model", "E2:E11")
    assert e["volatile"] and e["single_threaded_functions"] == ["INDIRECT"]
    assert group_at(out, "Model", "F2")["single_threaded_functions"] == ["CELL"]
    f3 = group_at(out, "Model", "F3")
    assert f3["volatile"] and not f3["single_threaded"]
    assert group_at(out, "Model", "G2")["single_threaded"]
    assert not group_at(out, "Model", "G3")["single_threaded"]
    h = group_at(out, "Model", "H2:H11")
    assert h["udfs"] == 1 and h["udf_names"] == ["MYUDF"] and h["single_threaded"]
    i = group_at(out, "Model", "I2:I11")
    assert i["array"] and i["cells"] == 10
    assert group_at(out, "Model", "J2")["whole_column_ref"]
    assert group_at(out, "Model", "K2")["dynamic_array"]
    assert group_at(out, "Model", "L2")["volatile_functions"] == ["NOW"]
    t = out["totals"]
    assert t["volatile_cells"] == 200 + 10 + 1 + 1 + 1 + 1
    assert t["udf_cells"] == 10
    assert t["single_thread_cells"] == 10 + 1 + 1 + 10
    assert t["array_cells"] == 10


def test_names_listing(synth_out):
    names = {n["name"]: n for n in synth_out["names"]}
    assert set(names) == {"InputBlock", "ModelBlock", "LocalOut", "GrowthRate", "DynBlock"}
    ib = names["InputBlock"]
    assert ib["is_range"] and ib["scope"] == "workbook" and ib["sheet"] == "Inputs"
    assert ib["refers_to_range"] == "Inputs!$A$2:$A$201" and ib["formula_cells"] == 0
    assert names["ModelBlock"]["formula_cells"] == 200
    lo = names["LocalOut"]
    assert lo["scope"] == "Report" and lo["is_range"]
    for n in ("GrowthRate", "DynBlock"):
        assert names[n]["is_range"] is False
        assert names[n]["refers_to_range"] is None and names[n]["sheet"] is None


def test_hashed_vs_clear(synth):
    hashed = inspect_workbook(synth, names="hashed", salt=SALT)
    clear = inspect_workbook(synth, names="clear")
    h = group_at(hashed, redact_name("Model", SALT), "H2:H11")
    assert h["udf_names"] == [redact_name("MYUDF", SALT)]
    assert hashed["redaction"] == {"names": "hashed"}
    assert {n["name"] for n in hashed["names"]} == {
        redact_name(x, SALT) for x in ("InputBlock", "ModelBlock", "LocalOut", "GrowthRate", "DynBlock")}
    ib = next(n for n in hashed["names"] if n["name"] == redact_name("InputBlock", SALT))
    assert ib["refers_to_range"] == redact_name("Inputs", SALT) + "!$A$2:$A$201"
    # built-in names are vocabulary and stay clear
    assert group_at(hashed, redact_name("Model", SALT), "D2:D201")["functions"] == ["OFFSET"]
    # hashing does not change grouping, only the fingerprint key
    assert [g["areas"] for g in hashed["groups"]] == [g["areas"] for g in clear["groups"]]
    with pytest.raises(ValueError):
        inspect_workbook(synth, names="hashed", salt="")


def test_plan_shape(synth):
    out, plan = inspect_workbook_and_plan(synth, names="hashed", salt=SALT)
    assert plan["in_memory_only"] is True
    model = next(s for s in plan["sheets"] if s["sheet"] == "Model")
    assert model["formula_area"] == "A2:L201" and model["array_areas"] == ["I2:I11"]
    assert {n["name"] for n in plan["names"]} == {"InputBlock", "ModelBlock", "LocalOut"}
    assert [g["group"] for g in plan["groups"]] == [g["group"] for g in out["groups"]]
    assert plan["redaction_map"][redact_name("Model", SALT)] == "Model"
    assert "redaction_map" not in inspect_workbook_for_plan(synth)


def test_write_formulas_json(tmp_path, synth):
    out, plan = inspect_workbook_and_plan(synth, names="hashed", salt=SALT)
    p = tmp_path / "formulas.json"
    write_formulas_json(out, p)
    assert json.loads(p.read_text()) == out
    with pytest.raises(ValueError):
        write_formulas_json(plan, tmp_path / "plan.json")


def test_xlsm_and_readonly(tmp_path, synth):
    before = synth.read_bytes()
    wb = load_workbook(synth)
    xlsm = tmp_path / "copy.xlsm"
    wb.save(xlsm)
    out = inspect_workbook(xlsm, names="clear")
    assert out["totals"]["formula_cells"] == inspect_workbook(synth, names="clear")["totals"]["formula_cells"]
    assert synth.read_bytes() == before


# --- privacy -----------------------------------------------------------------

SENTINEL_NUM = 918273.6455
SENTINEL_STR = "ZQXSENTINELZQX"


def test_privacy_no_values_or_formula_text(tmp_path):
    p = make_synthetic_workbook(tmp_path / "priv.xlsx", seed=11)
    wb = load_workbook(p)
    ws = wb["Model"]
    ws["N2"] = SENTINEL_NUM
    ws["N3"] = SENTINEL_STR
    ws["N4"] = '=IF(N2>%s,"%s",$A$1)' % (SENTINEL_NUM, SENTINEL_STR + "LIT")
    ws["N5"] = "=SecretFunc(N2)"
    wb.save(p)
    input_values = [c.value for c in wb["Inputs"]["A"] if isinstance(c.value, float)]

    for mode in ("hashed", "clear"):
        out = inspect_workbook(p, names=mode, salt=SALT)
        text = json.dumps(out)
        for needle in [str(SENTINEL_NUM), "918273", SENTINEL_STR, "SUM(", "($A$1", ",$A$1)", "$A$2:$A2",
                       "OFFSET(", "Inputs!A", "*2", '"address"', "RC[", "<N>", "<S>", "="]:
            assert needle not in text, (mode, needle)
        for v in input_values[:50]:
            assert repr(v) not in text, (mode, v)

    hashed = json.dumps(inspect_workbook(p, names="hashed", salt=SALT))
    for ident in ["Inputs", "Model", "Report", "MYUDF", "SecretFunc", "InputBlock",
                  "ModelBlock", "LocalOut", "GrowthRate", "DynBlock"]:
        assert ident not in hashed, ident


# --- review fixes --------------------------------------------------------------

def _patch_sheet(src, dst, sheet_data=None, dimension=None, extra_parts=None):
    """Copy an openpyxl-written xlsx, replacing sheet1's sheetData/dimension."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                s = data.decode()
                if sheet_data is not None:
                    s = re.sub(r"<sheetData>.*</sheetData>|<sheetData\s*/>",
                               "<sheetData>" + sheet_data + "</sheetData>", s, flags=re.S)
                if dimension is not None:
                    s = re.sub(r'<dimension ref="[^"]*"\s*/>', dimension, s)
                data = s.encode()
            zout.writestr(item, data)
        for name, text in (extra_parts or {}).items():
            zout.writestr(name, text)
    return dst


def _base(tmp_path):
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = 1
    p = tmp_path / "base.xlsx"
    wb.save(p)
    return p


def test_cell_info_type_splits_groups(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["B1"] = '=CELL("width",A1)'
    for r in range(2, 101):
        ws.cell(row=r, column=2, value='=CELL("address",A%d)' % r)
    p = tmp_path / "cell.xlsx"
    wb.save(p)
    out = inspect_workbook(p, names="clear")
    assert sorted((g["cells"], g["single_threaded"]) for g in out["groups"]) == [(1, False), (99, True)]
    assert out["totals"]["single_thread_cells"] == 99
    assert out["totals"]["volatile_cells"] == 100
    # the info_type is Excel vocabulary; arbitrary literals still collapse
    assert "<S:address>" in normalize_formula('=CELL("Address",A1)', 1, 2)
    assert fp('=CELL("zzsecret",A1)', 1, 2) == fp('=CELL("other",A1)', 1, 2)


def test_group_features_are_or_of_members():
    a = classify('=CELL("width",A1)')
    b = classify("=MYUDF(Other!A1:A2,B1)")
    m = a.merge(b)
    assert m.functions == {"CELL"} and m.udfs == {"MYUDF"}
    assert m.single_threaded == {"UDF"} and m.volatile == {"CELL"}
    assert m.cross_sheet and m.references == 2
    assert a.merge(a) is a


def test_let_lambda_parameters():
    f = classify("=_xlfn.LET(_xlpm.f,_xlfn.LAMBDA(_xlpm.x,_xlpm.x+1),_xlpm.f(A1))")
    assert f.functions == {"LET", "LAMBDA"}
    assert not f.udfs and not f.single_threaded
    assert f.lambda_calls == {"F"}
    assert f.references == 1
    assert fp("=_xlfn.LET(_xlpm.a,A1,_xlpm.a*2)", 1, 2) == fp("=_xlfn.LET(_xlpm.a,A2,_xlpm.a*2)", 2, 2)
    assert fp("=_xlfn.LET(_xlpm.a,A1,_xlpm.a*2)", 1, 2) != fp("=_xlfn.LET(_xlpm.b,A1,_xlpm.b*2)", 1, 2)


def test_prefixed_functions_trusted_as_builtin():
    f = classify("=_xlfn.FUTUREFUNC(A1)+_xlfn._xlws.OTHERNEW(1)+_xlfn.IMPORTCSV(\"x\")+USDOLLAR(A1)")
    assert f.functions == {"FUTUREFUNC", "OTHERNEW", "IMPORTCSV", "USDOLLAR"}
    assert not f.udfs
    assert {"IMPORTTEXT", "IMPORTCSV", "USDOLLAR"} <= BUILTIN_FUNCTIONS
    g = classify("=[1]!AddinFn(A1)")
    assert g.udfs == {"[1]!ADDINFN"} and g.external_ref


@pytest.mark.parametrize("dimension", ['<dimension ref="A1:A1"/>', "", None])
def test_wrong_dimension_tag_does_not_drop_formulas(tmp_path, dimension):
    rows = "".join('<row r="%d"><c r="B%d"><f>A%d*2</f><v>0</v></c></row>' % (r, r, r)
                   for r in range(1, 501))
    p = _patch_sheet(_base(tmp_path), tmp_path / "dim.xlsx", sheet_data=rows, dimension=dimension)
    out = inspect_workbook(p, names="clear")
    assert out["totals"]["formula_cells"] == 500
    assert out["groups"][0]["areas"] == ["B1:B500"]


_METADATA = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">
<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" pasteAll="1"
 pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1"
 coerce="1" cellMeta="1"/></metadataTypes>
<futureMetadata name="XLDAPR" count="1"><bk><extLst><ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}">
<xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/></ext></extLst></bk></futureMetadata>
<cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata>
</metadata>"""


def _spill_rows():
    rows = []
    for r in range(1, 101):
        cells = []
        if r == 1:
            cells.append('<c r="C1" cm="1"><f t="array" ref="C1:C100">_xlfn.SEQUENCE(100)</f><v>1</v></c>')
            cells.append('<c r="D1" cm="1"><f t="array" ref="D1:D100">A1:A100*2</f><v>1</v></c>')
            cells.append('<c r="E1"><f t="array" ref="E1:E10">A1:A10*3</f><v>1</v></c>')
        else:
            cells.append('<c r="C%d"><v>%d</v></c>' % (r, r))
            cells.append('<c r="D%d"><v>0</v></c>' % r)
        rows.append('<row r="%d">%s</row>' % (r, "".join(cells)))
    return "".join(rows)


@pytest.mark.parametrize("with_metadata", [True, False])
def test_spill_anchor_counts_once(tmp_path, with_metadata):
    extra = {"xl/metadata.xml": _METADATA} if with_metadata else None
    p = _patch_sheet(_base(tmp_path), tmp_path / "spill.xlsx", sheet_data=_spill_rows(),
                     extra_parts=extra)
    out, plan = inspect_workbook_and_plan(p, names="clear")
    c = group_at(out, "S", "C1")
    d = group_at(out, "S", "D1")
    e = group_at(out, "S", "E1:E10")
    assert (c["cells"], c["spill_cells"], c["array"], c["dynamic_array"]) == (1, 99, False, True)
    assert (d["cells"], d["spill_cells"], d["array"], d["dynamic_array"]) == (1, 99, False, True)
    assert (e["cells"], e["spill_cells"], e["array"], e["dynamic_array"]) == (10, 0, True, False)
    t = out["totals"]
    assert t["formula_cells"] == 12 and t["spill_cells"] == 198
    assert t["dynamic_array_cells"] == 2 and t["array_cells"] == 10
    assert out["sheets"][0]["spill_cells"] == 198
    s = plan["sheets"][0]
    assert s["spill_areas"] == ["C1:C100", "D1:D100"] and s["array_areas"] == ["E1:E10"]


def test_metadata_without_dynamic_flag_is_not_spill(tmp_path):
    meta = _METADATA.replace('fDynamic="1"', 'fDynamic="0"')
    p = _patch_sheet(_base(tmp_path), tmp_path / "nospill.xlsx", sheet_data=_spill_rows(),
                     extra_parts={"xl/metadata.xml": meta})
    out = inspect_workbook(p, names="clear")
    assert out["totals"]["spill_cells"] == 0
    assert group_at(out, "S", "D1:D100")["array"]


def test_name_counting_scales(tmp_path):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("S")
    for r in range(1, 20001):
        ws.append([r, "=A%d+$A$%d" % (r, r)])  # 20k single-cell groups
    for i in range(2000):
        wb.defined_names["N%d_x" % i] = DefinedName("N%d_x" % i, attr_text="S!$A$%d:$B$%d" % (i + 1, i + 50))
    p = tmp_path / "names.xlsx"
    wb.save(p)
    t0 = time.perf_counter()
    out, plan = inspect_workbook_and_plan(p, names="clear")
    elapsed = time.perf_counter() - t0
    assert out["totals"]["groups"] == 20000
    assert all(n["formula_cells"] == 50 for n in out["names"])
    assert all(n["formula_cells"] == 50 and n["has_formulas"] for n in plan["names"])
    assert elapsed < 12.0, elapsed  # was ~19 s with the per-name scan


def test_fingerprint_keyed_in_hashed_mode(tmp_path):
    wb = Workbook()
    wb.active.title = "Secret"
    wb.active["B2"] = "=SUM(AcquisitionTargetPrice)*[1]Deals!A1"
    p = tmp_path / "fp.xlsx"
    wb.save(p)
    norm = "=" + normalize_formula("=SUM(AcquisitionTargetPrice)*[1]Deals!A1", 2, 2)
    a = inspect_workbook(p, salt="salt-one")["groups"][0]["fingerprint"]
    b = inspect_workbook(p, salt="salt-two")["groups"][0]["fingerprint"]
    a2 = inspect_workbook(p, salt="salt-one")["groups"][0]["fingerprint"]
    clear = inspect_workbook(p, names="clear")["groups"][0]["fingerprint"]
    assert a == a2 and a != b
    assert clear == hashlib.sha256(norm.encode()).hexdigest()[:16] == fingerprint(norm)
    assert a == fingerprint(norm, "salt-one") and a != clear
    plan = inspect_workbook_for_plan(p, salt="salt-one")
    assert plan["groups"][0]["fingerprint"] == a


def test_sheet_prefix_canonical_and_udf_case(tmp_path):
    assert fp("='Sheet2'!A1", 1, 2) == fp("=Sheet2!A1", 1, 2) == fp("=sheet2!A1", 1, 2)
    assert fp("='It''s'!A1", 1, 2) == fp("='IT''S'!A1", 1, 2)
    assert fp("='Sheet2'!A1", 1, 2) != fp("=Sheet3!A1", 1, 2)
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["B1"] = "=MyUdf(A1)+'Other'!A1"
    ws["B2"] = "=MYUDF(A2)+Other!A2"
    wb.create_sheet("Other")
    p = tmp_path / "case.xlsx"
    wb.save(p)
    out = inspect_workbook(p, salt=SALT)
    (g,) = out["groups"]
    assert g["cells"] == 2 and g["udf_names"] == [redact_name("MYUDF", SALT)]
