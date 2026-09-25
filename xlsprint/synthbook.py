"""Synthetic workbook for tests and demos (openpyxl, generated numbers only).

The workbook exercises every feature the formula inspector detects. It holds
no real data: every constant is drawn from ``random.Random(seed)``.

Sheets
  Inputs  A2:B{n+1}  generated numbers
  Model   A  =Inputs!A2*Inputs!B2          copied-down relative group, cross-sheet
          B  =SUM($A$2:$A2)                period-to-date (Microsoft's slow pattern)
          C  =C1+A2 (C2 is =A2)            running-total alternative
          D  =OFFSET(A2,0,1)               volatile
          E  =INDIRECT("A"&ROW())          volatile + single-threaded (10 cells)
          F2 =CELL("address",A2)           volatile + single-threaded
          F3 =CELL("width",A2)             volatile only
          G2 =ADDRESS(1,1,1,TRUE,"Model")  single-threaded (sheet_text argument)
          G3 =ADDRESS(1,1)                 not single-threaded
          H  =MYUDF(A2)                    UDF call (10 cells)
          I2:I11 {=A2:A11*2}               CSE array formula
          J2 =SUM(Inputs!A:A)              whole-column reference
          K2 =_xlfn.SEQUENCE(5)            dynamic-array function
          L2 =NOW()  L3 =RAND()            volatile
  Report  A1:A3 cross-sheet summaries
Names
  InputBlock  (workbook)  Inputs!$A$2:$A${n+1}
  ModelBlock  (workbook)  Model!$B$2:$B${n+1}
  LocalOut    (Report)    Report!$A$1:$A$3
  GrowthRate  (workbook)  a generated constant   -> not a range
  DynBlock    (workbook)  OFFSET(...)            -> not a range
"""

from __future__ import annotations

import random
from pathlib import Path


def make_synthetic_workbook(path, *, seed: int = 0, rows: int = 200) -> Path:
    from openpyxl import Workbook
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.formula import ArrayFormula

    if rows < 12:
        raise ValueError("rows must be >= 12")
    rng = random.Random(seed)
    path = Path(path)
    last = rows + 1

    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    model = wb.create_sheet("Model")
    report = wb.create_sheet("Report")

    for r in range(2, last + 1):
        inputs.cell(row=r, column=1, value=round(rng.uniform(1, 1000), 2))
        inputs.cell(row=r, column=2, value=rng.randint(1, 50))

    for r in range(2, last + 1):
        model.cell(row=r, column=1, value="=Inputs!A%d*Inputs!B%d" % (r, r))
        model.cell(row=r, column=2, value="=SUM($A$2:$A%d)" % r)
        model.cell(row=r, column=3, value="=A2" if r == 2 else "=C%d+A%d" % (r - 1, r))
        model.cell(row=r, column=4, value="=OFFSET(A%d,0,1)" % r)
    for r in range(2, 12):
        model.cell(row=r, column=5, value='=INDIRECT("A"&ROW())')
        model.cell(row=r, column=8, value="=MYUDF(A%d)" % r)
    model["F2"] = '=CELL("address",A2)'
    model["F3"] = '=CELL("width",A2)'
    model["G2"] = '=ADDRESS(1,1,1,TRUE,"Model")'
    model["G3"] = "=ADDRESS(1,1)"
    model["I2"] = ArrayFormula("I2:I11", "=A2:A11*2")
    model["J2"] = "=SUM(Inputs!A:A)"
    model["K2"] = "=_xlfn.SEQUENCE(5)"
    model["L2"] = "=NOW()"
    model["L3"] = "=RAND()"

    report["A1"] = "=SUM(Model!C2:C%d)" % last
    report["A2"] = "=Model!B%d/COUNT(Inputs!A2:A%d)" % (last, last)
    report["A3"] = "=MAX(Model!A2:A%d)" % last

    def add_name(name, text, sheet=None):
        dn = DefinedName(name, attr_text=text)
        if sheet is None:
            wb.defined_names[name] = dn
        else:
            sheet.defined_names[name] = dn

    add_name("InputBlock", "Inputs!$A$2:$A$%d" % last)
    add_name("ModelBlock", "Model!$B$2:$B$%d" % last)
    add_name("LocalOut", "Report!$A$1:$A$3", sheet=report)
    add_name("GrowthRate", repr(round(rng.uniform(0.01, 0.09), 4)))
    add_name("DynBlock", "OFFSET(Inputs!$A$2,0,0,10,1)")

    wb.save(path)
    return path
