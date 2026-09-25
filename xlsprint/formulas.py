"""Static formula inspection and grouping (openpyxl, never evaluates).

``inspect_workbook`` produces the ``formulas.json`` object described in
docs/DESIGN.md. It is safe to write to disk: it carries counts, flags,
coordinates, built-in function names, and (by default) hashed identifiers,
but never formula text, literals, or cell values.

``inspect_workbook_for_plan`` returns the clear-text structure the runner needs
to drive Excel (real sheet names, areas, named ranges). It is in-memory only
and must not be written to disk by default.

Both come from one internal pass (``_scan``); ``inspect_workbook_and_plan``
returns both from a single read of the file.
"""

from __future__ import annotations

import bisect
import hashlib
import hmac
import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

SCHEMA = "xlsprint.formulas/1"
PLAN_SCHEMA = "xlsprint.plan/1"
MAX_LISTED_AREAS = 32

MAX_ROW = 1_048_576
MAX_COL = 16_384

LIMITATIONS = [
    "Static inspection only: no formula is evaluated and no timing is implied by these counts.",
    "UDFs declared volatile with Application.Volatile are not detected (the VBA project is not parsed in v1).",
    "Database functions (DSUM etc.) over PivotTables cannot be detected statically as single-threaded.",
    "CELL and ADDRESS are flagged single-threaded only when the argument check matches syntactically.",
    "Volatility or thread-safety inherited through defined names (e.g. a name that uses OFFSET) is not resolved.",
    "Conditional-format, data-validation, chart-series and defined-name formulas are not inspected as cell formulas.",
    "Data tables (what-if TABLE ranges) are counted but not grouped or analysed.",
    "A spilled dynamic-array formula counts as 1 formula cell at its anchor; the rest of the spill "
    "range is reported as spill_cells, not formula cells.",
    "Spill anchors are identified from the cell-metadata (cm) dynamic-array marker Excel writes; "
    "files written without it (older Excel, some generators) report such formulas as CSE array formulas.",
    "Excel 4 macro sheets, dialog sheets and chart sheets are not inspected.",
    "Group fingerprints hash the normalized formula structure with literals replaced by placeholders; "
    "in hashed mode they are keyed with the salt (HMAC) and cannot be compared across runs with different salts.",
    "Group features are the OR over all member cells; references is the maximum over members.",
]

# ---------------------------------------------------------------------------
# Function vocabularies
# ---------------------------------------------------------------------------

VOLATILE_FUNCTIONS = frozenset(
    "RAND RANDBETWEEN RANDARRAY NOW TODAY OFFSET INDIRECT CELL INFO".split()
)

# From Microsoft's article. CELL and ADDRESS are conditional (see _classify).
SINGLE_THREADED_FUNCTIONS = frozenset(
    """PHONETIC CELL INDIRECT GETPIVOTDATA CUBEMEMBER CUBEVALUE
    CUBEMEMBERPROPERTY CUBESET CUBERANKEDMEMBER CUBEKPIMEMBER CUBESETCOUNT
    ADDRESS ERROR.TYPE HYPERLINK""".split()
)

DYNAMIC_ARRAY_FUNCTIONS = frozenset(
    """ANCHORARRAY FILTER SORT SORTBY UNIQUE SEQUENCE RANDARRAY TOCOL TOROW
    WRAPCOLS WRAPROWS VSTACK HSTACK TAKE DROP EXPAND CHOOSECOLS CHOOSEROWS
    TEXTSPLIT MAKEARRAY MAP SCAN BYROW BYCOL GROUPBY PIVOTBY TRIMRANGE""".split()
)

BUILTIN_FUNCTIONS = frozenset(
    """
    ABS ACCRINT ACCRINTM ACOS ACOSH ACOT ACOTH ADDRESS AGGREGATE AMORDEGRC
    AMORLINC AND ANCHORARRAY ARABIC AREAS ARRAYTOTEXT ASC ASIN ASINH ATAN ATAN2
    ATANH AVEDEV AVERAGE AVERAGEA AVERAGEIF AVERAGEIFS BAHTTEXT BASE BESSELI
    BESSELJ BESSELK BESSELY BETA.DIST BETA.INV BETADIST BETAINV BIN2DEC BIN2HEX
    BIN2OCT BINOM.DIST BINOM.DIST.RANGE BINOM.INV BINOMDIST BITAND BITLSHIFT
    BITOR BITRSHIFT BITXOR BYCOL BYROW CALL CEILING CEILING.MATH
    CEILING.PRECISE CELL CHAR CHIDIST CHIINV CHISQ.DIST CHISQ.DIST.RT CHISQ.INV
    CHISQ.INV.RT CHISQ.TEST CHITEST CHOOSE CHOOSECOLS CHOOSEROWS CLEAN CODE
    COLUMN COLUMNS COMBIN COMBINA COMPLEX CONCAT CONCATENATE CONFIDENCE
    CONFIDENCE.NORM CONFIDENCE.T CONVERT COPILOT CORREL COS COSH COT COTH COUNT
    COUNTA COUNTBLANK COUNTIF COUNTIFS COUPDAYBS COUPDAYS COUPDAYSNC COUPNCD
    COUPNUM COUPPCD COVAR COVARIANCE.P COVARIANCE.S CRITBINOM CSC CSCH
    CUBEKPIMEMBER CUBEMEMBER CUBEMEMBERPROPERTY CUBERANKEDMEMBER CUBESET
    CUBESETCOUNT CUBEVALUE CUMIPMT CUMPRINC DATE DATEDIF DATEVALUE DAVERAGE DAY
    DAYS DAYS360 DB DBCS DCOUNT DCOUNTA DDB DEC2BIN DEC2HEX DEC2OCT DECIMAL
    DEGREES DELTA DETECTLANGUAGE DEVSQ DGET DISC DMAX DMIN DOLLAR DOLLARDE
    DOLLARFR DPRODUCT DROP DSTDEV DSTDEVP DSUM DURATION DVAR DVARP ECMA.CEILING
    EDATE EFFECT ENCODEURL EOMONTH ERF ERF.PRECISE ERFC ERFC.PRECISE ERROR.TYPE
    EUROCONVERT EVEN EXACT EXP EXPAND EXPON.DIST EXPONDIST F.DIST F.DIST.RT
    F.INV F.INV.RT F.TEST FACT FACTDOUBLE FALSE FDIST FIELDVALUE FILTER
    FILTERXML FIND FINDB FINV FISHER FISHERINV FIXED FLOOR FLOOR.MATH
    FLOOR.PRECISE FORECAST FORECAST.ETS FORECAST.ETS.CONFINT
    FORECAST.ETS.SEASONALITY FORECAST.ETS.STAT FORECAST.LINEAR FORMULATEXT
    FREQUENCY FTEST FV FVSCHEDULE GAMMA GAMMA.DIST GAMMA.INV GAMMADIST
    GAMMAINV GAMMALN GAMMALN.PRECISE GAUSS GCD GEOMEAN GESTEP GETPIVOTDATA
    GROUPBY GROWTH HARMEAN HEX2BIN HEX2DEC HEX2OCT HLOOKUP HOUR HSTACK
    HYPERLINK HYPGEOM.DIST HYPGEOMDIST IF IFERROR IFNA IFS IMABS IMAGE
    IMAGINARY IMARGUMENT IMCONJUGATE IMCOS IMCOSH IMCOT IMCSC IMCSCH IMDIV
    IMEXP IMLN IMLOG10 IMLOG2 IMPOWER IMPRODUCT IMREAL IMSEC IMSECH IMSIN
    IMSINH IMSQRT IMSUB IMSUM IMTAN INDEX INDIRECT INFO INT INTERCEPT INTRATE
    IPMT IRR ISBLANK ISERR ISERROR ISEVEN ISFORMULA ISLOGICAL ISNA ISNONTEXT
    ISNUMBER ISO.CEILING ISODD ISOMITTED ISOWEEKNUM ISPMT ISREF ISTEXT JIS
    KURT LAMBDA LARGE LCM LEFT LEFTB LEN LENB LET LINEST LN LOG LOG10 LOGEST
    LOGINV LOGNORM.DIST LOGNORM.INV LOGNORMDIST LOOKUP LOWER MAKEARRAY MAP
    MATCH MAX MAXA MAXIFS MDETERM MDURATION MEDIAN MID MIDB MIN MINA MINIFS
    MINUTE MINVERSE MIRR MMULT MOD MODE MODE.MULT MODE.SNGL MONTH MROUND
    MULTINOMIAL MUNIT N NA NEGBINOM.DIST NEGBINOMDIST NETWORKDAYS
    NETWORKDAYS.INTL NOMINAL NORM.DIST NORM.INV NORM.S.DIST NORM.S.INV
    NORMDIST NORMINV NORMSDIST NORMSINV NOT NOW NPER NPV NUMBERVALUE OCT2BIN
    OCT2DEC OCT2HEX ODD ODDFPRICE ODDFYIELD ODDLPRICE ODDLYIELD OFFSET OR
    PDURATION PEARSON PERCENTILE PERCENTILE.EXC PERCENTILE.INC PERCENTOF
    PERCENTRANK PERCENTRANK.EXC PERCENTRANK.INC PERMUT PERMUTATIONA PHI
    PHONETIC PI PIVOTBY PMT POISSON POISSON.DIST POWER PPMT PRICE PRICEDISC
    PRICEMAT PROB PRODUCT PROPER PV PY QUARTILE QUARTILE.EXC QUARTILE.INC
    QUOTIENT RADIANS RAND RANDARRAY RANDBETWEEN RANK RANK.AVG RANK.EQ RATE
    RECEIVED REDUCE REGEXEXTRACT REGEXREPLACE REGEXTEST REGISTER.ID REPLACE
    REPLACEB REPT RIGHT RIGHTB ROMAN ROUND ROUNDDOWN ROUNDUP ROW ROWS RRI RSQ
    RTD SCAN SEARCH SEARCHB SEC SECH SECOND SEQUENCE SERIESSUM SHEET SHEETS
    SIGN SIN SINGLE SINH SKEW SKEW.P SLN SLOPE SMALL SORT SORTBY SQL.REQUEST
    SQRT SQRTPI STANDARDIZE STDEV STDEV.P STDEV.S STDEVA STDEVP STDEVPA STEYX
    STOCKHISTORY SUBSTITUTE SUBTOTAL SUM SUMIF SUMIFS SUMPRODUCT SUMSQ
    SUMX2MY2 SUMX2PY2 SUMXMY2 SWITCH SYD T T.DIST T.DIST.2T T.DIST.RT T.INV
    T.INV.2T T.TEST TABLE TAKE TAN TANH TBILLEQ TBILLPRICE TBILLYIELD TDIST
    TEXT TEXTAFTER TEXTBEFORE TEXTJOIN TEXTSPLIT TIME TIMEVALUE TINV TOCOL
    TODAY TOROW TRANSLATE TRANSPOSE TREND TRIM TRIMMEAN TRIMRANGE TRUE TRUNC
    TTEST TYPE UNICHAR IMPORTTEXT IMPORTCSV USDOLLAR UNICODE UNIQUE UPPER VALUE VALUETOTEXT VAR VAR.P VAR.S
    VARA VARP VARPA VDB VLOOKUP VSTACK WEBSERVICE WEEKDAY WEEKNUM WEIBULL
    WEIBULL.DIST WORKDAY WORKDAY.INTL WRAPCOLS WRAPROWS XIRR XLOOKUP XMATCH
    XNPV XOR YEAR YEARFRAC YIELD YIELDDISC YIELDMAT Z.TEST ZTEST
    """.split()
)

# Prefixes Excel writes in the file format. _xludf/_xll mark non-built-ins.
_BUILTIN_PREFIXES = ("_XLFN._XLWS.", "_XLFN.", "_XLWS.")
_UDF_PREFIXES = ("_XLUDF.", "_XLL.")
_PARAM_PREFIX = "_XLPM."

_ERRORS = (
    "#GETTING_DATA", "#CONNECT!", "#BLOCKED!", "#UNKNOWN!", "#PYTHON!",
    "#DIV/0!", "#VALUE!", "#FIELD!", "#SPILL!", "#CALC!", "#BUSY!", "#NULL!",
    "#NAME?", "#REF!", "#NUM!", "#N/A",
)

# ---------------------------------------------------------------------------
# Tokenizer
# ---------------------------------------------------------------------------

# Token kinds
OPERAND_REF = "REF"      # cell/range/name/structured reference (text as written)
OPERAND_NUM = "NUM"
OPERAND_STR = "STR"
OPERAND_BOOL = "BOOL"
OPERAND_ERR = "ERR"
FUNC_OPEN = "FUNC"       # value = function name as written (without "(")
PAREN_OPEN = "("
CLOSE = ")"
ARRAY_OPEN = "{"
ARRAY_CLOSE = "}"
SEP_ARG = ","
SEP_ROW = ";"
OP = "OP"
WS = "WS"

_OPERAND_KINDS = frozenset([OPERAND_REF, OPERAND_NUM, OPERAND_STR, OPERAND_BOOL, OPERAND_ERR])
_NUMBER_RE = re.compile(r"^(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?$")
_EXP_PENDING_RE = re.compile(r"^(?:\d+\.?\d*|\.\d+)[eE]$")
_OPERATOR_CHARS = "+-*/^&=<>%@:"


def _is_operand_char(ch: str) -> bool:
    return ch.isalnum() or ch in "_.$:!\\?"


def tokenize(formula: str) -> List[Tuple[str, str]]:
    """Split a formula into (kind, text) tokens. Never raises.

    Tolerant by design: unknown characters become OP tokens so every formula
    produces a stable token stream.
    """
    s = formula[1:] if formula.startswith("=") else formula
    n = len(s)
    i = 0
    out: List[Tuple[str, str]] = []
    # stack of "F" (function), "P" (paren), "A" (array constant)
    stack: List[str] = []

    def prev_significant() -> Optional[Tuple[str, str]]:
        for t in reversed(out):
            if t[0] != WS:
                return t
        return None

    while i < n:
        ch = s[i]
        if ch.isspace():
            j = i
            while j < n and s[j].isspace():
                j += 1
            out.append((WS, " "))
            i = j
        elif ch == '"':
            j = i + 1
            while j < n:
                if s[j] == '"':
                    if j + 1 < n and s[j + 1] == '"':
                        j += 2
                        continue
                    break
                j += 1
            out.append((OPERAND_STR, s[i:j + 1]))
            i = j + 1
        elif ch == "#":
            prev = prev_significant()
            upper = s[i:i + 14].upper()
            err = next((e for e in _ERRORS if upper.startswith(e)), None)
            if err is not None and not (prev and prev[0] == OPERAND_REF and prev[1].endswith("!")):
                out.append((OPERAND_ERR, err))
                i += len(err)
            elif err is not None:
                # Sheet1!#REF!
                out[-1] = (OPERAND_REF, out[-1][1] + err)
                i += len(err)
            else:
                out.append((OP, "#"))  # spill operator A1#
                i += 1
        elif ch == "{":
            stack.append("A")
            out.append((ARRAY_OPEN, "{"))
            i += 1
        elif ch == "}":
            if stack and stack[-1] == "A":
                stack.pop()
            out.append((ARRAY_CLOSE, "}"))
            i += 1
        elif ch == "(":
            stack.append("P")
            out.append((PAREN_OPEN, "("))
            i += 1
        elif ch == ")":
            if stack and stack[-1] in "FP":
                stack.pop()
            out.append((CLOSE, ")"))
            i += 1
        elif ch == ",":
            if stack and stack[-1] in "FA":
                out.append((SEP_ARG, ","))
            else:
                out.append((OP, ","))  # union operator
            i += 1
        elif ch == ";":
            out.append((SEP_ROW, ";"))
            i += 1
        elif ch == "'" or ch == "[" or _is_operand_char(ch):
            j = i
            while j < n:
                c = s[j]
                if c == "'":
                    k = j + 1
                    while k < n:
                        if s[k] == "'":
                            if k + 1 < n and s[k + 1] == "'":
                                k += 2
                                continue
                            break
                        k += 1
                    j = k + 1
                elif c == "[":
                    depth = 0
                    k = j
                    while k < n:
                        if s[k] == "[":
                            depth += 1
                        elif s[k] == "]":
                            depth -= 1
                            if depth == 0:
                                break
                        elif s[k] == "'" and k + 1 < n:
                            k += 1  # escaped char inside structured ref
                        k += 1
                    j = k + 1
                elif _is_operand_char(c):
                    j += 1
                else:
                    break
            text = s[i:j]
            # 1.5E+3: pull in the exponent sign and digits
            if j < n and s[j] in "+-" and _EXP_PENDING_RE.match(text):
                k = j + 1
                while k < n and s[k].isdigit():
                    k += 1
                text = s[i:k]
                j = k
            i = j
            if i < n and s[i] == "(":
                # function call; a range prefix like A1:INDEX( is split off
                name = text
                if ":" in name and "!" not in name.rsplit(":", 1)[1]:
                    head, name = name.rsplit(":", 1)
                    out.append((OPERAND_REF, head))
                    out.append((OP, ":"))
                stack.append("F")
                out.append((FUNC_OPEN, name))
                i += 1
            elif _NUMBER_RE.match(text):
                out.append((OPERAND_NUM, text))
            elif text.upper() in ("TRUE", "FALSE"):
                out.append((OPERAND_BOOL, text.upper()))
            else:
                out.append((OPERAND_REF, text))
        elif ch in "<>" and i + 1 < n and s[i + 1] in "=>" and s[i:i + 2] in ("<=", ">=", "<>"):
            out.append((OP, s[i:i + 2]))
            i += 2
        else:
            out.append((OP, ch))
            i += 1
    return out


# ---------------------------------------------------------------------------
# Reference parsing
# ---------------------------------------------------------------------------

_CELL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d{1,7})$")
_COL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})$")
_ROW_RE = re.compile(r"^(\$?)(\d{1,7})$")


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def _col_letters(idx: int) -> str:
    out = ""
    while idx:
        idx, rem = divmod(idx - 1, 26)
        out = chr(65 + rem) + out
    return out


@dataclass
class RefInfo:
    kind: str                       # "cell" | "area" | "cols" | "rows" | "name" | "error"
    prefix: Optional[str] = None    # sheet/workbook prefix as written (without "!")
    sheet: Optional[str] = None     # unquoted sheet name when single-sheet prefix
    external: bool = False
    three_d: bool = False
    book: Optional[str] = None      # external workbook part (index or path), unquoted
    sheet_text: Optional[str] = None  # unquoted sheet part (may be "S1:S3" for 3-D)
    # parts: list of (row_abs, row, col_abs, col); row/col None when absent
    parts: List[Tuple[bool, Optional[int], bool, Optional[int]]] = field(default_factory=list)
    name: Optional[str] = None


def _split_prefix(text: str) -> Tuple[Optional[str], str]:
    """Split 'prefix!ref' at the last '!' outside quotes/brackets."""
    in_q = False
    depth = 0
    last = -1
    i = 0
    while i < len(text):
        c = text[i]
        if in_q:
            if c == "'":
                if i + 1 < len(text) and text[i + 1] == "'":
                    i += 2
                    continue
                in_q = False
        elif c == "'":
            in_q = True
        elif c == "[":
            depth += 1
        elif c == "]":
            depth -= 1
        elif c == "!" and depth == 0:
            last = i
        i += 1
    if last < 0:
        return None, text
    return text[:last], text[last + 1:]


def _parse_part(part: str):
    m = _CELL_RE.match(part)
    if m:
        col, row = _col_index(m.group(2)), int(m.group(4))
        if 1 <= col <= MAX_COL and 1 <= row <= MAX_ROW:
            return ("cell", (bool(m.group(3)), row, bool(m.group(1)), col))
    m = _COL_RE.match(part)
    if m:
        col = _col_index(m.group(2))
        if 1 <= col <= MAX_COL:
            return ("col", (False, None, bool(m.group(1)), col))
    m = _ROW_RE.match(part)
    if m:
        row = int(m.group(2))
        if 1 <= row <= MAX_ROW:
            return ("row", (bool(m.group(1)), row, False, None))
    return None


def parse_ref(text: str) -> RefInfo:
    if text.upper().endswith("!#REF!"):
        prefix, body = _split_prefix(text[:-5])
        body = "#REF!"
    else:
        prefix, body = _split_prefix(text)
    info = RefInfo(kind="name", prefix=prefix)
    if prefix is not None:
        p = prefix
        if p.startswith("'") and p.endswith("'") and len(p) >= 2:
            p = p[1:-1].replace("''", "'")
        if "[" in p or p == "":
            info.external = True
        m = re.match(r"^(.*)\[([^\]]*)\](.*)$", p)
        if m:
            info.book = m.group(1) + "[" + m.group(2) + "]"
            p = m.group(3)
        elif p == "":
            info.book = ""
        info.sheet_text = p
        if ":" in p:
            info.three_d = True
        else:
            info.sheet = p or None
    if "[" in body or body.upper().endswith("#REF!") or body == "":
        info.kind = "error" if body.upper().endswith("#REF!") else "name"
        info.name = body
        return info
    pieces = body.split(":")
    parsed = [_parse_part(p) for p in pieces]
    if all(parsed) and len(pieces) <= 2:
        kinds = {p[0] for p in parsed}
        if kinds == {"cell"}:
            info.kind = "cell" if len(pieces) == 1 else "area"
            info.parts = [p[1] for p in parsed]
            return info
        if kinds == {"col"} and len(pieces) == 2:
            info.kind = "cols"
            info.parts = [p[1] for p in parsed]
            return info
        if kinds == {"row"} and len(pieces) == 2:
            info.kind = "rows"
            info.parts = [p[1] for p in parsed]
            return info
    info.kind = "name"
    info.name = body
    return info


def _r1c1_part(part, row: int, col: int) -> str:
    r_abs, r, c_abs, c = part
    out = ""
    if r is not None:
        if r_abs:
            out += "R%d" % r
        else:
            d = r - row
            out += "R" if d == 0 else "R[%d]" % d
    if c is not None:
        if c_abs:
            out += "C%d" % c
        else:
            d = c - col
            out += "C" if d == 0 else "C[%d]" % d
    return out


# ---------------------------------------------------------------------------
# Normalization and classification
# ---------------------------------------------------------------------------

def _clean_func_name(name: str) -> Tuple[str, str]:
    """Return (name without file-format prefix, tag).

    tag: "fn" for _xlfn./_xlws. (trusted built-in), "udf" for _xludf./_xll.
    or external [n]! calls, "param" for _xlpm. (LET/LAMBDA parameter), "".
    """
    up = name.upper()
    if up.startswith(_PARAM_PREFIX):
        return name[len(_PARAM_PREFIX):], "param"
    for p in _UDF_PREFIXES:
        if up.startswith(p):
            return name[len(p):], "udf"
    for p in _BUILTIN_PREFIXES:
        if up.startswith(p):
            return name[len(p):], "fn"
    if "!" in name:
        return name, "udf"
    return name, ""


# CELL info_type values: kept in the normalized form (Excel vocabulary),
# because "format"/"address" change the single-threaded classification.
_CELL_INFO_TYPES = frozenset(
    "address col color contents filename format parentheses prefix protect row type width".split()
)


def _cell_info_literal(sig: Sequence[Tuple[str, str]], i: int) -> Optional[str]:
    """If sig[i] is the sole string literal first argument of CELL, its info_type."""
    if i == 0 or i + 1 >= len(sig) or sig[i][0] != OPERAND_STR:
        return None
    prev, nxt = sig[i - 1], sig[i + 1]
    if prev[0] != FUNC_OPEN or nxt[0] not in (SEP_ARG, CLOSE):
        return None
    name, tag = _clean_func_name(prev[1])
    if tag not in ("", "fn") or name.upper() != "CELL":
        return None
    value = sig[i][1][1:-1].strip().lower()
    return value if value in _CELL_INFO_TYPES else "?"


def _normalize_tokens(tokens: Sequence[Tuple[str, str]], row: int, col: int,
                      host_sheet: Optional[str] = None) -> str:
    parts: List[str] = []
    prev_kind = None
    sig = [t for t in tokens if t[0] != WS]
    si = -1
    for idx, (kind, text) in enumerate(tokens):
        if kind == WS:
            nxt = next((t for t in tokens[idx + 1:] if t[0] != WS), None)
            if (prev_kind in _OPERAND_KINDS or prev_kind in (CLOSE,)) and nxt is not None and (
                    nxt[0] in _OPERAND_KINDS or nxt[0] in (FUNC_OPEN, PAREN_OPEN)):
                parts.append(" ")  # intersection operator
            continue
        prev_kind = kind
        si += 1
        if kind == OPERAND_NUM:
            parts.append("<N>")
        elif kind == OPERAND_STR:
            info_type = _cell_info_literal(sig, si)
            parts.append("<S>" if info_type is None else "<S:%s>" % info_type)
        elif kind == OPERAND_BOOL or kind == OPERAND_ERR:
            parts.append(text.upper())
        elif kind == FUNC_OPEN:
            name, tag = _clean_func_name(text)
            marker = {"udf": "U:", "param": "P:"}.get(tag, "")
            parts.append(marker + name.upper() + "(")
        elif kind == OPERAND_REF:
            parts.append(_normalize_ref(text, row, col, host_sheet))
        else:
            parts.append(text)
    return "".join(parts)


def _normalize_ref(text: str, row: int, col: int, host_sheet: Optional[str]) -> str:
    if text.upper().startswith(_PARAM_PREFIX):
        return "P:" + text[len(_PARAM_PREFIX):].upper()
    info = parse_ref(text)
    prefix = ""
    if info.prefix is not None:
        same = (info.sheet is not None and not info.external and host_sheet is not None
                and info.sheet.upper() == host_sheet.upper())
        if not same:
            # canonical: unquoted, unescaped, case-folded
            book = (info.book or "").upper() if info.external else ""
            prefix = "S:" + book + (info.sheet_text or "").upper() + "!"
    if info.kind in ("cell", "area", "cols", "rows"):
        return prefix + ":".join(_r1c1_part(p, row, col) for p in info.parts)
    return prefix + "N:" + (info.name or "").upper()


def normalize_formula(formula: str, row: int, col: int) -> str:
    """Relative-R1C1 normalized form of ``formula`` hosted at (row, col).

    Internal: the result is hashed for the fingerprint and never output.
    """
    return _normalize_tokens(tokenize(formula), row, col)


def fingerprint(normalized: str, key: Optional[str] = None) -> str:
    """sha256(normalized)[:16]; HMAC-SHA256(key, normalized)[:16] when keyed."""
    data = normalized.encode("utf-8")
    if key:
        return hmac.new(key.encode("utf-8"), data, hashlib.sha256).hexdigest()[:16]
    return hashlib.sha256(data).hexdigest()[:16]


def length_bucket(n: int) -> str:
    if n < 64:
        return "<64"
    if n < 256:
        return "64-255"
    if n < 1024:
        return "256-1023"
    return ">=1024"


@dataclass
class Features:
    functions: frozenset          # built-in names (clean, upper)
    udfs: frozenset               # non-built-in names (clean, upper)
    lambda_calls: frozenset       # calls to defined names / LAMBDA parameters (upper)
    volatile: frozenset
    single_threaded: frozenset    # built-in names, plus "UDF" pseudo-entry
    dynamic_array: bool
    cross_sheet: bool
    external_ref: bool
    whole_column_ref: bool
    whole_row_ref: bool
    references: int

    def merge(self, other: "Features") -> "Features":
        """OR of two feature sets (a group's features cover every member)."""
        if other is self or other == self:
            return self
        return Features(
            self.functions | other.functions, self.udfs | other.udfs,
            self.lambda_calls | other.lambda_calls, self.volatile | other.volatile,
            self.single_threaded | other.single_threaded,
            self.dynamic_array or other.dynamic_array, self.cross_sheet or other.cross_sheet,
            self.external_ref or other.external_ref,
            self.whole_column_ref or other.whole_column_ref,
            self.whole_row_ref or other.whole_row_ref,
            max(self.references, other.references))


def _classify(tokens: Sequence[Tuple[str, str]], host_sheet: Optional[str],
              defined_names: frozenset = frozenset()) -> Features:
    functions = set()
    udfs = set()
    lambda_calls = set()
    single = set()
    dynamic = False
    cross = external = whole_col = whole_row = False
    refs = 0

    sig = [t for t in tokens if t[0] != WS]
    # frame: [kind, clean_name, arg_index, first_arg_tokens, arg_has_content]
    frames: List[list] = []
    for kind, text in sig:
        if frames and frames[-1][0] == "F" and kind not in (SEP_ARG, CLOSE):
            fr = frames[-1]
            fr[4] = True
            if fr[2] == 0:
                fr[3].append((kind, text))
        if kind == FUNC_OPEN:
            clean, tag = _clean_func_name(text)
            up = clean.upper()
            if tag == "param" or (tag == "" and up in defined_names):
                lambda_calls.add(up)
            elif tag == "fn" or (tag == "" and up in BUILTIN_FUNCTIONS):
                functions.add(up)
                if up in DYNAMIC_ARRAY_FUNCTIONS:
                    dynamic = True
            else:
                udfs.add(up)
                if "!" in clean:
                    external = True
            frames.append(["F", up if tag in ("", "fn") else "", 0, [], False])
        elif kind in (PAREN_OPEN, ARRAY_OPEN):
            frames.append(["P", "", 0, [], False])
        elif kind in (CLOSE, ARRAY_CLOSE):
            if frames:
                fr = frames.pop()
                if fr[0] == "F":
                    _check_conditional(fr, single)
        elif kind == SEP_ARG:
            if frames and frames[-1][0] == "F":
                frames[-1][2] += 1
                frames[-1][4] = False
        elif kind == OP and text == "#":
            dynamic = True
        elif kind == OPERAND_REF:
            if text.upper().startswith(_PARAM_PREFIX):
                continue  # LET/LAMBDA parameter, not a reference
            refs += 1
            info = parse_ref(text)
            if info.external:
                external = True
            if info.three_d:
                cross = True
            elif info.sheet is not None and not info.external:
                if host_sheet is None or info.sheet.upper() != host_sheet.upper():
                    cross = True
            if info.kind == "cols":
                whole_col = True
            elif info.kind == "rows":
                whole_row = True
    # unterminated calls (malformed formulas): still evaluate what we saw
    for fr in frames:
        if fr[0] == "F":
            _check_conditional(fr, single)

    for f in functions:
        if f in SINGLE_THREADED_FUNCTIONS and f not in ("CELL", "ADDRESS"):
            single.add(f)
    if udfs:
        single.add("UDF")
    volatile = frozenset(f for f in functions if f in VOLATILE_FUNCTIONS)
    return Features(frozenset(functions), frozenset(udfs), frozenset(lambda_calls),
                    volatile, frozenset(single), dynamic, cross, external, whole_col,
                    whole_row, refs)


def _check_conditional(frame: list, single: set) -> None:
    _, name, arg_idx, first, last_has = frame
    if name == "CELL":
        if len(first) == 1 and first[0][0] == OPERAND_STR:
            info_type = first[0][1].strip('"').strip().lower()
            if info_type in ("format", "address"):
                single.add("CELL")
    elif name == "ADDRESS":
        # sheet_text is the 5th argument; flag only when it is present and non-empty
        if arg_idx > 4 or (arg_idx == 4 and last_has):
            single.add("ADDRESS")


# ---------------------------------------------------------------------------
# Area compaction
# ---------------------------------------------------------------------------

Rect = Tuple[int, int, int, int]  # r1, c1, r2, c2


def _runs(sorted_vals: Iterable[int]) -> List[Tuple[int, int]]:
    runs: List[Tuple[int, int]] = []
    for v in sorted_vals:
        if runs and v == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], v)
        elif not runs or v != runs[-1][1]:
            runs.append((v, v))
    return runs


def _merge(lines: Dict[int, List[int]], transpose: bool) -> List[Rect]:
    # lines: major index -> minor indices. Build runs per major, then merge
    # identical runs across consecutive majors.
    segs = []
    for major in sorted(lines):
        for a, b in _runs(sorted(lines[major])):
            segs.append((a, b, major))
    segs.sort()
    rects: List[Rect] = []
    cur = None
    for a, b, major in segs:
        if cur and cur[0] == a and cur[1] == b and cur[3] == major - 1:
            cur[3] = major
        else:
            if cur:
                rects.append(tuple(cur))
            cur = [a, b, major, major]
    if cur:
        rects.append(tuple(cur))
    out = []
    for a, b, m1, m2 in rects:
        if transpose:  # major = row, minor = col
            out.append((m1, a, m2, b))
        else:          # major = col, minor = row
            out.append((a, m1, b, m2))
    return out


def compact_areas(cells: Iterable[Tuple[int, int]]) -> List[Rect]:
    """Cover a set of (row, col) cells with rectangles; fewest of two sweeps."""
    by_col: Dict[int, List[int]] = {}
    by_row: Dict[int, List[int]] = {}
    for r, c in cells:
        by_col.setdefault(c, []).append(r)
        by_row.setdefault(r, []).append(c)
    a = _merge(by_col, transpose=False)
    b = _merge(by_row, transpose=True)
    best = a if len(a) <= len(b) else b
    return sorted(best)


def rect_a1(rect: Rect) -> str:
    r1, c1, r2, c2 = rect
    first = "%s%d" % (_col_letters(c1), r1)
    if (r1, c1) == (r2, c2):
        return first
    return "%s:%s%d" % (first, _col_letters(c2), r2)


def _a1_to_rect(ref: str) -> Optional[Rect]:
    info = parse_ref(ref)
    if info.kind == "cell":
        _, r, _, c = info.parts[0]
        return (r, c, r, c)
    if info.kind == "area" and len(info.parts) == 2:
        (_, r1, _, c1), (_, r2, _, c2) = info.parts
        return (min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))
    return None


# ---------------------------------------------------------------------------
# Redaction
# ---------------------------------------------------------------------------

def redact_name(name: str, salt: str) -> str:
    """Same rule as trace.redact_name: "h:" + sha256(salt+name)[:10]."""
    return "h:" + hashlib.sha256((salt + name).encode("utf-8")).hexdigest()[:10]


class _Namer:
    def __init__(self, mode: str, salt: Optional[str]):
        if mode not in ("hashed", "clear"):
            raise ValueError("names must be 'hashed' or 'clear'")
        if mode == "hashed" and not salt:
            raise ValueError("a non-empty salt is required when names='hashed'")
        self.mode = mode
        self.salt = salt or ""
        self.mapping: Dict[str, str] = {}

    def __call__(self, name: str) -> str:
        if self.mode == "clear":
            return name
        h = redact_name(name, self.salt)
        self.mapping[h] = name
        return h


# ---------------------------------------------------------------------------
# Internal scan
# ---------------------------------------------------------------------------

@dataclass
class _Group:
    sheet_index: int
    norm: str                    # normalized form (grouping key; never output)
    array: bool                  # CSE array formula
    spill: bool                  # dynamic-array spill anchor
    features: Features
    cols: Dict[int, List[int]] = field(default_factory=dict)
    cells: int = 0
    spill_cells: int = 0
    first: Tuple[int, int] = (MAX_ROW + 1, MAX_COL + 1)
    max_len: int = 0

    def add(self, row: int, col: int) -> None:
        self.cols.setdefault(col, []).append(row)
        self.cells += 1
        if (row, col) < self.first:
            self.first = (row, col)


@dataclass
class _Sheet:
    index: int
    title: str
    formula_cells: int = 0
    spill_cells: int = 0
    used: Optional[List[int]] = None     # r1, c1, r2, c2
    fbox: Optional[List[int]] = None
    data_tables: List[Rect] = field(default_factory=list)
    array_areas: List[Rect] = field(default_factory=list)
    spill_areas: List[Rect] = field(default_factory=list)


@dataclass
class _Name:
    name: str
    scope_sheet: Optional[str]   # None = workbook
    hidden: bool
    sheet: Optional[str]
    rect: Optional[Rect]
    address: Optional[str]       # absolute A1 address, coordinates only
    is_range: bool
    formula_cells: int = 0


@dataclass
class _Scan:
    sha_prefix: str
    sheets: List[_Sheet]
    groups: List[_Group]
    names: List[_Name]


def _expand(box: Optional[List[int]], r1: int, c1: int, r2: int, c2: int) -> List[int]:
    if box is None:
        return [r1, c1, r2, c2]
    box[0] = min(box[0], r1)
    box[1] = min(box[1], c1)
    box[2] = max(box[2], r2)
    box[3] = max(box[3], c2)
    return box


def _file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _collect_names(wb) -> List[Tuple[object, Optional[str]]]:
    out = []
    try:
        items = list(wb.defined_names.items())
    except AttributeError:  # pragma: no cover - openpyxl < 3.1
        items = [(d.name, d) for d in wb.defined_names.definedName]
    for name, defn in items:
        out.append((defn, None))
    for ws in wb.worksheets:
        for name, defn in getattr(ws, "defined_names", {}).items():
            out.append((defn, ws.title))
    return out


def _parse_name(defn, scope: Optional[str], sheet_titles: Dict[str, str]) -> Optional[_Name]:
    name = defn.name
    if name.upper().startswith("_XLNM."):
        return None
    hidden = bool(getattr(defn, "hidden", False))
    text = defn.attr_text or ""
    toks = [t for t in tokenize("=" + text) if t[0] != WS]
    rect = sheet = address = None
    is_range = False
    if len(toks) == 1 and toks[0][0] == OPERAND_REF:
        info = parse_ref(toks[0][1])
        if (info.sheet is not None and not info.external and not info.three_d
                and info.sheet.upper() in sheet_titles):
            sheet = sheet_titles[info.sheet.upper()]
            if info.kind in ("cell", "area") and len(info.parts) <= 2:
                ps = info.parts
                r1, c1 = ps[0][1], ps[0][3]
                r2, c2 = ps[-1][1], ps[-1][3]
                rect = (min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))
                address = "$%s$%d" % (_col_letters(rect[1]), rect[0])
                if len(ps) == 2:
                    address += ":$%s$%d" % (_col_letters(rect[3]), rect[2])
                is_range = True
            elif info.kind == "cols":
                c1, c2 = sorted(p[3] for p in info.parts)
                rect = (1, c1, MAX_ROW, c2)
                address = "$%s:$%s" % (_col_letters(c1), _col_letters(c2))
                is_range = True
            elif info.kind == "rows":
                r1, r2 = sorted(p[1] for p in info.parts)
                rect = (r1, 1, r2, MAX_COL)
                address = "$%d:$%d" % (r1, r2)
                is_range = True
            if not is_range:
                sheet = None
    return _Name(name=name, scope_sheet=scope, hidden=hidden, sheet=sheet, rect=rect,
                 address=address, is_range=is_range)


# --- dynamic arrays ----------------------------------------------------------
#
# Excel stores a spilled dynamic-array formula as an array formula (t="array")
# on the anchor cell, marked with a cell-metadata index (c/@cm) that points to
# an XLDAPR dynamicArrayProperties record with fDynamic="1" in xl/metadata.xml.
# openpyxl drops @cm, so it is read here straight from the sheet XML, and only
# for sheets that contain array formulas.

_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _dynamic_cm_indices(archive) -> Optional[frozenset]:
    """1-based cell-metadata indices marking dynamic arrays; None if unknown."""
    import xml.etree.ElementTree as ET

    try:
        names = set(archive.namelist())
        path = "xl/metadata.xml"
        if path not in names:
            return None
        root = ET.fromstring(archive.read(path))
        types = [mt.get("name") for mt in root.iter(_MAIN_NS + "metadataType")]
        future: Dict[str, list] = {}
        for fm in root.findall(_MAIN_NS + "futureMetadata"):
            future[fm.get("name")] = fm.findall(_MAIN_NS + "bk")
        dyn = set()
        cm_block = root.find(_MAIN_NS + "cellMetadata")
        if cm_block is None:
            return frozenset()
        for i, bk in enumerate(cm_block.findall(_MAIN_NS + "bk"), start=1):
            for rc in bk.findall(_MAIN_NS + "rc"):
                t, v = int(rc.get("t", "0")), int(rc.get("v", "0"))
                if not 1 <= t <= len(types) or types[t - 1] != "XLDAPR":
                    continue
                recs = future.get("XLDAPR", [])
                if 0 <= v < len(recs):
                    for el in recs[v].iter():
                        if (_local(el.tag) == "dynamicArrayProperties"
                                and el.get("fDynamic") in ("1", "true")):
                            dyn.add(i)
        return frozenset(dyn)
    except Exception:
        return None


def _spill_anchors(archive, sheet_path: Optional[str], dyn_cm: Optional[frozenset]) -> set:
    """Coordinates (row, col) of array-formula anchors flagged as dynamic arrays.

    With dyn_cm None (no readable metadata part) any @cm on an array anchor
    counts; Excel uses cell metadata only for dynamic arrays.
    """
    import xml.etree.ElementTree as ET

    out = set()
    if archive is None or not sheet_path or dyn_cm == frozenset():
        return out
    try:
        with archive.open(sheet_path) as src:
            for _, el in ET.iterparse(src, events=("end",)):
                if el.tag != _MAIN_NS + "c":
                    continue
                cm = el.get("cm")
                if cm is not None and (dyn_cm is None or (cm.isdigit() and int(cm) in dyn_cm)):
                    f = el.find(_MAIN_NS + "f")
                    if f is not None and f.get("t") == "array":
                        rect = _a1_to_rect(el.get("r", ""))
                        if rect:
                            out.add((rect[0], rect[1]))
                el.clear()
    except Exception:
        return set()
    return out


def _scan(path) -> _Scan:
    from openpyxl import load_workbook
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    path = Path(path)
    sha = _file_sha256(path)
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False,
                       keep_vba=False, rich_text=False)
    try:
        archive = getattr(wb, "_archive", None)
        dyn_cm = _dynamic_cm_indices(archive) if archive is not None else None
        raw_names = _collect_names(wb)
        defined_upper = frozenset(d.name.upper() for d, _ in raw_names)
        sheet_titles = {ws.title.upper(): ws.title for ws in wb.worksheets}
        sheets: List[_Sheet] = []
        groups: Dict[Tuple[int, str], _Group] = {}
        feature_cache: Dict[Tuple[str, str], Tuple[list, Features]] = {}

        def add_formula(sh, text, r, c, rect, kind):
            # kind: "" plain, "array" CSE, "spill" dynamic-array anchor
            key = (sh.title, text)
            cached = feature_cache.get(key)
            if cached is None:
                toks = tokenize(text)
                cached = (toks, _classify(toks, sh.title, defined_upper))
                if len(feature_cache) < 200_000:
                    feature_cache[key] = cached
            toks, feats = cached
            lead = {"array": "{=", "spill": "#="}.get(kind, "=")
            norm = lead + _normalize_tokens(toks, r, c, sh.title)
            g = groups.get((sh.index, norm))
            if g is None:
                g = groups[(sh.index, norm)] = _Group(
                    sheet_index=sh.index, norm=norm, array=kind == "array",
                    spill=kind == "spill", features=feats)
            else:
                g.features = g.features.merge(feats)
            g.max_len = max(g.max_len, len(text))
            if kind == "array":
                for rr in range(rect[0], rect[2] + 1):
                    for cc in range(rect[1], rect[3] + 1):
                        g.add(rr, cc)
                sh.array_areas.append(rect)
                sh.formula_cells += (rect[2] - rect[0] + 1) * (rect[3] - rect[1] + 1)
                sh.fbox = _expand(sh.fbox, *rect)
            else:
                g.add(r, c)
                sh.formula_cells += 1
                sh.fbox = _expand(sh.fbox, r, c, r, c)
                if kind == "spill":
                    extra = (rect[2] - rect[0] + 1) * (rect[3] - rect[1] + 1) - 1
                    g.spill_cells += extra
                    sh.spill_cells += extra
                    sh.spill_areas.append(rect)
            sh.used = _expand(sh.used, *rect)

        for s_idx, ws in enumerate(wb.worksheets):
            sh = _Sheet(index=s_idx, title=ws.title)
            sheets.append(sh)
            # The <dimension> tag can be wrong or stale; read_only mode would
            # silently truncate to it.
            reset = getattr(ws, "reset_dimensions", None)
            if reset is not None:
                reset()
            arrays = []
            for row in ws.iter_rows():
                for cell in row:
                    value = getattr(cell, "value", None)
                    if value is None:
                        continue
                    r, c = cell.row, cell.column
                    sh.used = _expand(sh.used, r, c, r, c)
                    if isinstance(value, DataTableFormula):
                        rect = _a1_to_rect(value.ref or "") or (r, c, r, c)
                        sh.data_tables.append(rect)
                        sh.used = _expand(sh.used, *rect)
                    elif isinstance(value, ArrayFormula):
                        rect = _a1_to_rect(value.ref or "") or (r, c, r, c)
                        arrays.append((value.text or "=", r, c, rect))
                    elif cell.data_type == "f" and isinstance(value, str):
                        add_formula(sh, value, r, c, (r, c, r, c), "")
            if arrays:
                anchors = _spill_anchors(archive, getattr(ws, "_worksheet_path", None), dyn_cm)
                for text, r, c, rect in arrays:
                    add_formula(sh, text, r, c, rect, "spill" if (r, c) in anchors else "array")

        names = []
        for defn, scope in raw_names:
            parsed = _parse_name(defn, scope, sheet_titles)
            if parsed is not None:
                names.append(parsed)
    finally:
        close = getattr(wb, "close", None)
        if close:
            close()

    ordered = sorted(groups.values(), key=lambda g: (g.sheet_index, g.first, g.norm))
    _count_name_cells(sheets, ordered, names)
    return _Scan(sha_prefix=sha[:10], sheets=sheets, groups=ordered, names=names)


def _count_name_cells(sheets: List[_Sheet], groups: List[_Group], names: List[_Name]) -> None:
    """Formula cells inside each range name, via one per-sheet column index."""
    wanted = {nm.sheet for nm in names if nm.is_range}
    index: Dict[str, Tuple[List[int], Dict[int, List[int]]]] = {}
    by_idx = {s.index: s.title for s in sheets}
    for g in groups:
        title = by_idx[g.sheet_index]
        if title not in wanted:
            continue
        cols = index.setdefault(title, ([], {}))[1]
        for c, rows in g.cols.items():
            cols.setdefault(c, []).extend(rows)
    for title, (keys, cols) in index.items():
        for rows in cols.values():
            rows.sort()
        keys.extend(sorted(cols))
    for nm in names:
        if not nm.is_range or nm.sheet not in index:
            continue
        keys, cols = index[nm.sheet]
        r1, c1, r2, c2 = nm.rect
        total = 0
        for k in keys[bisect.bisect_left(keys, c1):bisect.bisect_right(keys, c2)]:
            rows = cols[k]
            total += bisect.bisect_right(rows, r2) - bisect.bisect_left(rows, r1)
        nm.formula_cells = total


def _group_rects(g: _Group) -> List[Rect]:
    cells = ((r, c) for c, rows in g.cols.items() for r in rows)
    return compact_areas(cells)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def _build_output(scan: _Scan, namer: _Namer, rects: Dict[int, List[Rect]],
                  fps: List[str]) -> dict:
    sheet_out = []
    for sh in scan.sheets:
        sheet_out.append({
            "sheet": namer(sh.title),
            "formula_cells": sh.formula_cells,
            "spill_cells": sh.spill_cells,
            "used_range": rect_a1(tuple(sh.used)) if sh.used else None,
            "data_tables": len(sh.data_tables),
        })

    totals = {"formula_cells": 0, "spill_cells": 0, "groups": len(scan.groups),
              "volatile_cells": 0, "udf_cells": 0, "single_thread_cells": 0, "array_cells": 0,
              "dynamic_array_cells": 0, "cross_sheet_cells": 0, "external_ref_cells": 0,
              "whole_column_ref_cells": 0, "whole_row_ref_cells": 0,
              "data_tables": sum(len(s.data_tables) for s in scan.sheets)}
    group_out = []
    for i, g in enumerate(scan.groups):
        f = g.features
        rs = rects[i]
        title = scan.sheets[g.sheet_index].title
        entry = {
            "group": "G%04d" % (i + 1),
            "fingerprint": fps[i],
            "sheet": namer(title),
            "cells": g.cells,
            "spill_cells": g.spill_cells,
            "areas": [rect_a1(r) for r in rs[:MAX_LISTED_AREAS]],
            "areas_total": len(rs),
            "functions": sorted(f.functions),
            "udfs": len(f.udfs),
            "udf_names": sorted(namer(u) for u in f.udfs),
            "lambda_calls": len(f.lambda_calls),
            "volatile": bool(f.volatile),
            "volatile_functions": sorted(f.volatile),
            "single_threaded": bool(f.single_threaded),
            "single_threaded_functions": sorted(f.single_threaded),
            "array": g.array,
            "dynamic_array": f.dynamic_array or g.spill,
            "cross_sheet": f.cross_sheet,
            "external_ref": f.external_ref,
            "whole_column_ref": f.whole_column_ref,
            "whole_row_ref": f.whole_row_ref,
            "references": f.references,
            "length_bucket": length_bucket(g.max_len),
        }
        group_out.append(entry)
        totals["formula_cells"] += g.cells
        totals["spill_cells"] += g.spill_cells
        for flag, key in ((entry["volatile"], "volatile_cells"),
                          (bool(f.udfs), "udf_cells"),
                          (entry["single_threaded"], "single_thread_cells"),
                          (g.array, "array_cells"),
                          (entry["dynamic_array"], "dynamic_array_cells"),
                          (f.cross_sheet, "cross_sheet_cells"),
                          (f.external_ref, "external_ref_cells"),
                          (f.whole_column_ref, "whole_column_ref_cells"),
                          (f.whole_row_ref, "whole_row_ref_cells")):
            if flag:
                totals[key] += g.cells

    name_out = []
    for nm in scan.names:
        entry = {
            "name": namer(nm.name),
            "scope": "workbook" if nm.scope_sheet is None else namer(nm.scope_sheet),
            "refers_to_range": None,
            "sheet": None,
            "is_range": nm.is_range,
            "hidden": nm.hidden,
        }
        if nm.is_range:
            entry["sheet"] = namer(nm.sheet)
            entry["refers_to_range"] = entry["sheet"] + "!" + nm.address
            entry["formula_cells"] = nm.formula_cells
        name_out.append(entry)
    totals["names"] = len(name_out)
    totals["range_names"] = sum(1 for n in name_out if n["is_range"])

    return {
        "schema": SCHEMA,
        "workbook_sha256_prefix": scan.sha_prefix,
        "redaction": {"names": namer.mode},
        "sheets": sheet_out,
        "groups": group_out,
        "names": name_out,
        "totals": totals,
        "limitations": list(LIMITATIONS),
    }


def _quote_sheet(title: str) -> str:
    if re.match(r"^[A-Za-z_][A-Za-z0-9_.]*$", title) and not _CELL_RE.match(title):
        return title
    return "'" + title.replace("'", "''") + "'"


def _build_plan(scan: _Scan, rects: Dict[int, List[Rect]], fps: List[str],
                namer: Optional[_Namer]) -> dict:
    plan_sheets = []
    for sh in scan.sheets:
        plan_sheets.append({
            "sheet": sh.title,
            "index": sh.index,
            "formula_cells": sh.formula_cells,
            "formula_area": rect_a1(tuple(sh.fbox)) if sh.fbox else None,
            "used_range": rect_a1(tuple(sh.used)) if sh.used else None,
            "array_areas": [rect_a1(r) for r in sh.array_areas],
            "spill_areas": [rect_a1(r) for r in sh.spill_areas],
        })
    plan_names = []
    for nm in scan.names:
        if not nm.is_range:
            continue
        plan_names.append({
            "name": nm.name,
            "scope": "workbook" if nm.scope_sheet is None else nm.scope_sheet,
            "sheet": nm.sheet,
            "address": nm.address,
            "refers_to_range": _quote_sheet(nm.sheet) + "!" + nm.address,
            "is_range": True,
            "formula_cells": nm.formula_cells,
            "has_formulas": nm.formula_cells > 0,
        })
    plan_groups = []
    for i, g in enumerate(scan.groups):
        rs = rects[i]
        plan_groups.append({
            "group": "G%04d" % (i + 1),
            "fingerprint": fps[i],
            "sheet": scan.sheets[g.sheet_index].title,
            "cells": g.cells,
            "areas": [rect_a1(r) for r in rs],
            "areas_total": len(rs),
            "array": g.array,
        })
    plan = {
        "schema": PLAN_SCHEMA,
        "in_memory_only": True,
        "workbook_sha256_prefix": scan.sha_prefix,
        "sheets": plan_sheets,
        "names": plan_names,
        "groups": plan_groups,
    }
    if namer is not None and namer.mode == "hashed":
        # clear identifiers the output hashed: lets the runner redact trace
        # names consistently. Local only.
        plan["redaction_map"] = dict(namer.mapping)
    return plan


def inspect_workbook_and_plan(path, *, names: str = "hashed", salt: Optional[str] = None
                              ) -> Tuple[dict, dict]:
    """One pass: (formulas.json object, in-memory plan). See the two wrappers.

    Fingerprints are HMAC-SHA256(salt, normalized)[:16] when names="hashed"
    and sha256(normalized)[:16] when names="clear"; the plan carries the same
    values as the formulas.json object.
    """
    namer = _Namer(names, salt)
    scan = _scan(path)
    rects = {i: _group_rects(g) for i, g in enumerate(scan.groups)}
    key = namer.salt if namer.mode == "hashed" else None
    fps = [fingerprint(g.norm, key) for g in scan.groups]
    out = _build_output(scan, namer, rects, fps)
    plan = _build_plan(scan, rects, fps, namer)
    return out, plan


def inspect_workbook(path, *, names: str = "hashed", salt: Optional[str] = None) -> dict:
    """The ``formulas.json`` object. Contains no formula text, literals or values."""
    return inspect_workbook_and_plan(path, names=names, salt=salt)[0]


def inspect_workbook_for_plan(path, *, salt: Optional[str] = None) -> dict:
    """Clear-text structure for driving Excel. In-memory only; do not persist.

    With ``salt`` it also carries ``redaction_map`` (hashed -> clear) for the
    identifiers inspect_workbook would hash with that salt, and its group
    fingerprints match inspect_workbook(path, names="hashed", salt=salt).
    """
    names = "hashed" if salt else "clear"
    return inspect_workbook_and_plan(path, names=names, salt=salt)[1]


def write_formulas_json(obj, path) -> None:
    if obj.get("schema") != SCHEMA:
        raise ValueError("not a formulas.json object (plans are in-memory only)")
    path = Path(path)
    tmp = path.with_name(path.name + ".tmp")
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, indent=1, sort_keys=False)
        fh.write("\n")
    os.replace(tmp, path)
